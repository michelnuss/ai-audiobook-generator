"""Read the source .docx into speaker-tagged paragraphs and write the English .docx.

Speaker identity lives in the highlight color of each paragraph's runs
(YELLOW = Ilana, TURQUOISE = Baruck, BRIGHT_GREEN = ET), the same mapping used by
app/parse_colors.py. Inline bold inside a paragraph is carried through the
model as **markup** so it can be restored on the English runs.
"""

import copy
import re
import shutil
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.paragraph import Paragraph
from docx.text.run import Run

HIGHLIGHT_TO_SPEAKER = {
    WD_COLOR_INDEX.YELLOW: "ilana",
    WD_COLOR_INDEX.TURQUOISE: "baruck",
    WD_COLOR_INDEX.BRIGHT_GREEN: "et",
}
UNMARKED_SPEAKER = "unmarked"

ROMAN_NUMERAL = re.compile(r"^M{0,4}(CM|CD|D?C{0,3})(XC|XL|L?X{0,3})(IX|IV|V?I{0,3})$")


@dataclass
class SourceParagraph:
    id: str
    index: int
    text: str  # plain text, or text with **bold** markup when bold is mixed
    speaker: str
    translatable: bool


@dataclass
class BuildReport:
    translated: int = 0
    kept_source: list[str] = field(default_factory=list)
    markup_dropped: list[str] = field(default_factory=list)


def paragraph_id(index: int) -> str:
    return f"P{index:04d}"


def is_roman_numeral(text: str) -> bool:
    return bool(text) and bool(ROMAN_NUMERAL.match(text.strip()))


def _text_runs(paragraph: Paragraph) -> list[Run]:
    return [r for r in paragraph.runs if r.text]


def dominant_speaker(paragraph: Paragraph) -> str:
    counts: Counter = Counter()
    for run in _text_runs(paragraph):
        counts[HIGHLIGHT_TO_SPEAKER.get(run.font.highlight_color, UNMARKED_SPEAKER)] += len(run.text)
    return counts.most_common(1)[0][0] if counts else UNMARKED_SPEAKER


def has_mixed_bold(paragraph: Paragraph) -> bool:
    runs = [r for r in _text_runs(paragraph) if r.text.strip()]
    return len({bool(r.bold) for r in runs}) > 1


def render_markup(paragraph: Paragraph) -> str:
    """Paragraph text; bold spans are wrapped in ** only when bold is mixed."""
    if not has_mixed_bold(paragraph):
        return paragraph.text
    parts: list[str] = []
    in_bold = False
    for run in _text_runs(paragraph):
        bold = bool(run.bold)
        if bold != in_bold:
            parts.append("**")
            in_bold = bold
        parts.append(run.text)
    if in_bold:
        parts.append("**")
    return "".join(parts).replace("****", "")


def parse_markup(text: str) -> list[tuple[str, bool]] | None:
    """Split '**bold** plain' into (segment, is_bold) pairs; None if markers are unbalanced."""
    parts = text.split("**")
    if len(parts) % 2 == 0:
        return None
    return [(part, i % 2 == 1) for i, part in enumerate(parts) if part]


def read_source(path: Path) -> list[SourceParagraph]:
    doc = Document(str(path))
    result = []
    for index, para in enumerate(doc.paragraphs):
        text = render_markup(para)
        plain = para.text.strip()
        result.append(
            SourceParagraph(
                id=paragraph_id(index),
                index=index,
                text=text.strip(),
                speaker=dominant_speaker(para),
                translatable=bool(plain) and not is_roman_numeral(plain),
            )
        )
    return result


def _insert_run_like(template: Run, paragraph: Paragraph, text: str, bold: bool | None) -> None:
    """Insert a run before `template` with the same run properties (highlight, font, size)."""
    new_r = copy.deepcopy(template._r)
    for child in list(new_r):
        if child is not new_r.rPr:
            new_r.remove(child)
    template._r.addprevious(new_r)
    run = Run(new_r, paragraph)
    run.text = text
    if bold is not None:
        run.bold = bold


def replace_paragraph_text(paragraph: Paragraph, text: str) -> bool:
    """Replace the paragraph's text, keeping the first run's formatting (highlight included).

    Returns False when the text had unbalanced ** markup and was written as plain text.
    """
    runs = _text_runs(paragraph)
    if not runs:
        paragraph.add_run(text.replace("**", ""))
        return "**" not in text

    mixed = has_mixed_bold(paragraph)
    markup_ok = True
    if mixed or "**" in text:
        parsed = parse_markup(text)
        if parsed is None:
            markup_ok = False
            segments = [(text.replace("**", ""), None)]
        elif mixed:
            segments = [(seg, bold) for seg, bold in parsed]
        else:
            segments = [(seg, True if bold else None) for seg, bold in parsed]
    else:
        segments = [(text, None)]

    template = runs[0]
    for segment, bold in segments:
        _insert_run_like(template, paragraph, segment, bold)
    for run in runs:
        run._r.getparent().remove(run._r)
    return markup_ok


def write_translation(source: Path, output: Path, translations: dict[str, str]) -> BuildReport:
    """Build the English document from a copy of the source, paragraph by paragraph."""
    output.parent.mkdir(parents=True, exist_ok=True)
    tmp = output.with_suffix(".tmp.docx")
    shutil.copyfile(source, tmp)
    doc = Document(str(tmp))
    report = BuildReport()
    for index, para in enumerate(doc.paragraphs):
        pid = paragraph_id(index)
        plain = para.text.strip()
        if not plain or is_roman_numeral(plain):
            continue
        english = translations.get(pid)
        if english is None:
            report.kept_source.append(pid)
            continue
        if not replace_paragraph_text(para, english):
            report.markup_dropped.append(pid)
        report.translated += 1
    doc.save(str(tmp))
    tmp.replace(output)
    return report


def highlight_signature(path: Path) -> list[str]:
    """Per-paragraph speaker sequence, used to verify the output against the source."""
    return [dominant_speaker(p) if p.text.strip() else "" for p in Document(str(path)).paragraphs]


def verify_output(source: Path, output: Path) -> list[str]:
    """Return a list of problems; empty means paragraph count and speakers match."""
    src, out = highlight_signature(source), highlight_signature(output)
    problems = []
    if len(src) != len(out):
        problems.append(f"paragraph count differs: source {len(src)}, output {len(out)}")
    for i, (a, b) in enumerate(zip(src, out)):
        if a != b:
            problems.append(f"{paragraph_id(i)}: speaker {a!r} became {b!r}")
    return problems
