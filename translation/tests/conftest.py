import json
import re
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import prompts  # noqa: E402
from config import Settings  # noqa: E402
from llm import LLMResult  # noqa: E402

Y, T, G = WD_COLOR_INDEX.YELLOW, WD_COLOR_INDEX.TURQUOISE, WD_COLOR_INDEX.BRIGHT_GREEN

# (runs as (text, highlight, bold)) — P0000..P0009
FIXTURE = [
    [("Aviso ao leitor.", Y, False)],
    [],
    [("Capítulo 1: Introdução de Baruck", T, True)],
    [("I", Y, True)],
    [("O fluido do perispírito sustenta o corpo.", G, False)],
    [("Baruck", Y, True), (", meu Mestre, falou.", Y, False)],
    [("Capítulo 2: Sintonia", T, True)],
    [("A sintonia com os irmãos estelares.", G, False)],
    [("Ele desencarnou em paz.", T, False)],
    [("Posfácio: O Despertar", Y, True)],
]

SECRET_REASONING = "SECRET-TRANSLATOR-REASONING"


def make_docx(path: Path) -> Path:
    doc = Document()
    for runs in FIXTURE:
        para = doc.add_paragraph()
        for text, highlight, bold in runs:
            run = para.add_run(text)
            run.font.highlight_color = highlight
            run.bold = bold
    doc.save(str(path))
    return path


@pytest.fixture
def source_docx(tmp_path) -> Path:
    return make_docx(tmp_path / "source.docx")


@pytest.fixture
def settings(tmp_path, source_docx) -> Settings:
    return Settings(
        source_docx=source_docx,
        output_docx=tmp_path / "output" / "book-en.docx",
        state_dir=tmp_path / "state",
        review_dir=tmp_path / "review",
        model="test-model",
        effort="high",
        review_effort="xhigh",
        max_chunk_words=12,
        context_paragraphs=40,
        max_output_tokens=1000,
    )


def _block(user: str, tag: str):
    match = re.search(rf"<{tag}>\n(.*?)\n</{tag}>", user, re.S)
    return json.loads(match.group(1)) if match else None


class FakeLLM:
    """Answers each agent by schema; records every call for assertions."""

    def __init__(self, critic_issue: bool = True, fail_on: str | None = None, refuse_on: str | None = None):
        self.calls: list[dict] = []
        self.critic_issue = critic_issue
        self.fail_on = fail_on
        self.refuse_on = refuse_on

    def kinds(self) -> list[str]:
        return [c["kind"] for c in self.calls]

    def complete_json(self, *, system, user, schema, effort, max_tokens) -> LLMResult:
        kind = {id(prompts.BIBLE_SCHEMA): "bible", id(prompts.CRITIQUE_SCHEMA): "critic",
                id(prompts.BOOK_REVIEW_SCHEMA): "review"}.get(id(schema))
        if kind is None:
            kind = "refine" if "<current_translation>" in user else "translate"
        self.calls.append({"kind": kind, "system": system, "user": user, "effort": effort})
        if self.fail_on == kind:
            raise RuntimeError(f"simulated crash in {kind}")
        if self.refuse_on == kind:
            from llm import RefusalError
            raise RefusalError("simulated refusal")
        usage = {"input_tokens": 10, "output_tokens": 5, "cache_read_input_tokens": 100,
                 "cache_creation_input_tokens": 0}
        return LLMResult(data=getattr(self, kind)(user), usage=usage)

    def bible(self, user):
        return {"style_guide": "Dignified.",
                "characters": [{"speaker": s, "description": s, "voice": "calm"} for s in ("ilana", "baruck", "et")],
                "terms": [{"source": "irmãos estelares", "target": "star brothers", "rationale": "recurring"}]}

    def translate(self, user):
        chunk = _block(user, "source_chunk")
        return {"translations": [{"id": p["id"], "text": "EN:" + p["text"]} for p in chunk],
                "flags": [{"paragraph_id": chunk[0]["id"], "original_term": "sintonia",
                           "assigned_translation": "attunement", "context_snippet": "A sintonia",
                           "ai_reasoning": SECRET_REASONING}]}

    def refine(self, user):
        current = _block(user, "current_translation")
        return {"translations": [{"id": p["id"], "text": "REFINED:" + p["text"]} for p in current],
                "flags": []}

    def critic(self, user):
        chunk = _block(user, "source_chunk")
        issues = []
        if self.critic_issue:
            issues = [{"paragraph_id": chunk[0]["id"], "severity": "major", "category": "tone",
                       "problem": "too casual", "suggested_fix": "more dignified"}]
        return {"issues": issues, "flags": []}

    def review(self, user):
        decisions = _block(user, "decisions")
        return {
            "overall_verdict": "fit_after_revisions",
            "report": "Solid translation.",
            "scores": [{"dimension": "doctrinal_fidelity", "score": 9, "commentary": "good"}],
            "chapters": [{"chunk_id": c["chunk_id"], "verdict": "fit", "notes": ""} for c in _block(user, "chunks")],
            "decision_verdicts": [{"decision_id": d["decision_id"], "verdict": "disagree",
                                   "preferred_translation": "attunement (sintonia)", "rationale": "context"}
                                  for d in decisions],
            "findings": [{"paragraph_ids": ["P0008"], "category": "secularized_concept", "severity": "critical",
                          "problem": "'desencarnou' rendered as died", "suggested_revision": "He discarnated in peace."}],
        }
