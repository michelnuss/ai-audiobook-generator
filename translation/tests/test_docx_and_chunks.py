import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from agents import TranslationValidationError, validate_translations
from chunker import build_chunks
from docx_io import parse_markup, read_source, verify_output, write_translation
from glossary import Glossary, GlossaryEntry


def test_read_source_speakers_markup_and_translatable(source_docx):
    paras = {p.id: p for p in read_source(source_docx)}
    assert paras["P0000"].speaker == "ilana"
    assert paras["P0002"].speaker == "baruck"
    assert paras["P0004"].speaker == "et"
    assert not paras["P0001"].translatable  # empty
    assert not paras["P0003"].translatable  # roman numeral
    assert paras["P0005"].text == "**Baruck**, meu Mestre, falou."
    assert paras["P0004"].text == "O fluido do perispírito sustenta o corpo."  # no markup when uniform


def test_write_translation_keeps_highlight_bold_and_structure(source_docx, tmp_path):
    out = tmp_path / "en.docx"
    translations = {p.id: f"EN {p.id}" for p in read_source(source_docx) if p.translatable}
    translations["P0005"] = "**Baruck**, my Master, spoke."
    report = write_translation(source_docx, out, translations)

    assert report.translated == 8 and not report.kept_source and not report.markup_dropped
    assert verify_output(source_docx, out) == []
    paras = Document(str(out)).paragraphs
    assert paras[3].text == "I"  # roman numeral untouched
    runs = [r for r in paras[5].runs if r.text]
    assert [(r.text, bool(r.bold)) for r in runs] == [("Baruck", True), (", my Master, spoke.", False)]
    assert all(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in runs)
    heading = [r for r in paras[2].runs if r.text]
    assert heading[0].bold and heading[0].font.highlight_color == WD_COLOR_INDEX.TURQUOISE
    assert paras[8].runs[0].font.highlight_color == WD_COLOR_INDEX.TURQUOISE


def test_unbalanced_markup_falls_back_to_plain(source_docx, tmp_path):
    out = tmp_path / "en.docx"
    report = write_translation(source_docx, out, {"P0005": "**Baruck, my Master, spoke."})
    assert report.markup_dropped == ["P0005"]
    assert Document(str(out)).paragraphs[5].text == "Baruck, my Master, spoke."
    assert "P0000" in report.kept_source


def test_parse_markup():
    assert parse_markup("a **b** c") == [("a ", False), ("b", True), (" c", False)]
    assert parse_markup("a **b") is None


def test_chunks_follow_chapters_and_size_limit(source_docx):
    paras = read_source(source_docx)
    chunks = build_chunks(paras, max_words=12)
    all_ids = [pid for c in chunks for pid in c.paragraph_ids]
    assert all_ids == [p.id for p in paras if p.translatable]
    # chapter headings always start a chunk
    starts = {c.paragraph_ids[0] for c in chunks}
    assert {"P0002", "P0006"} <= starts
    big = build_chunks(paras, max_words=10_000)
    assert len(big) == 1


def test_validate_translations():
    ok = validate_translations(["P1", "P2"], {"translations": [{"id": "P1", "text": "a"}, {"id": "P2", "text": "b"}]})
    assert ok == {"P1": "a", "P2": "b"}
    with pytest.raises(TranslationValidationError, match="missing paragraph ids: P2"):
        validate_translations(["P1", "P2"], {"translations": [{"id": "P1", "text": "a"}]})
    with pytest.raises(TranslationValidationError, match="unknown"):
        validate_translations(["P1"], {"translations": [{"id": "P1", "text": "a"}, {"id": "X", "text": "b"}]})
    with pytest.raises(TranslationValidationError, match="empty"):
        validate_translations(["P1"], {"translations": [{"id": "P1", "text": " "}]})


def test_glossary_precedence():
    g = Glossary.seeded()
    assert not g.add(GlossaryEntry("Perispírito", "astral body", origin="bible"))
    assert g.add(GlossaryEntry("perispirito", "perispirit (reviewed)", origin="reviewer"))
    assert "perispirit (reviewed)" in g.render()
    assert "HUMAN REVIEWER DECISION" in g.render()


def test_foundry_client_ignores_empty_endpoint_setting(monkeypatch):
    from llm import FoundryLLM
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_API_KEY", "test-key")
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_RESOURCE", "")
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_BASE_URL", "https://example.services.ai.azure.com/anthropic")
    client = FoundryLLM("m").client
    assert str(client.base_url).startswith("https://example.services.ai.azure.com/anthropic")
    monkeypatch.setenv("ANTHROPIC_FOUNDRY_RESOURCE", "example")
    FoundryLLM("m")  # both filled in: no error
